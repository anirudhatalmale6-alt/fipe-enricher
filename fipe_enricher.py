"""
FIPE Table Enricher - Hybrid offline/API approach
==================================================
Cars: matched against local CSV dataset (fipe_completa.csv)
Motos: matched against cached API data + live API fallback

Usage:
    python fipe_enricher.py lotes_leilao.xlsx

Requirements:
    pip install openpyxl requests
"""

import sys
import time
import json
import re
import os
import csv
from difflib import SequenceMatcher
from collections import defaultdict
import requests
import openpyxl

API_BASE = "https://parallelum.com.br/fipe/api/v1"
CACHE_FILE = "fipe_cache.json"
CSV_FILE = "fipe_completa.csv"

BRAND_MAP_CARROS = {
    "CHEVROLET": "GM - Chevrolet", "CHEV": "GM - Chevrolet", "GM": "GM - Chevrolet",
    "FIAT": "Fiat", "FORD": "Ford", "VW": "VW - VolksWagen",
    "PEUGEOT": "Peugeot", "RENAULT": "Renault", "CITROEN": "Citroën",
    "M.BENZ": "Mercedes-Benz", "MERCEDES": "Mercedes-Benz",
    "TOYOTA": "Toyota", "HYUNDAI": "Hyundai", "HONDA": "Honda",
    "CHERY": "Caoa Chery/Chery", "GEELY": "GEELY",
}

BRAND_MAP_MOTOS_API = {
    "YAMAHA": "101", "SUNDOWN": "98", "DAFRA": "145",
    "SHINERAY": "134", "JTZ": "209", "HONDA": "80",
}

MOTO_KEYWORDS = [
    "CG", "CBX", "BIZ", "NXR", "XRE", "CB ", "TITAN", "FAN", "FACTOR",
    "FAZER", "YBR", "NEO", "LANDER", "CRYPTON", "LEAD", "PCX", "POP",
    "BROS", "TWISTER", "TORNADO", "FALCON", "SPEED", "PHOENIX", "CHOPPER",
    "CARGO", "START", "CROSSER", "XTZ", "FZ25", "BURGMAN", "INTRUDER",
    "C100", "C 100", "MAX", "WEB",
]

CAR_KEYWORDS = [
    "CIVIC", "FIT", "CITY", "HR-V", "HRV", "WR-V", "ACCORD", "CR-V",
    "ML", "CLASS", "TUCSON", "CRETA", "HB20", "IX35",
]

MODEL_ALIASES = {
    "MEGANESD": "MEGANE SEDAN",
    "207HB": "207 HB",
    "206HB": "206 HB",
    "307HB": "307 HB",
    "SCENICEXP": "SCENIC EXP",
    "EC-7": "EC7",
}

NOISE_WORDS = {
    "gl", "gls", "glx", "ex", "exs", "lx", "lt", "ltz", "ls", "se", "sx",
    "dx", "xs", "xr", "ghia", "attract", "comfort", "expression", "dynamique",
    "privilege", "prestige", "premier", "at", "mt", "aut", "mec", "flex",
    "flexone", "flexpower", "econoflex", "mpfi", "mpi", "efi", "vhc",
    "8v", "16v", "24v", "32v", "2p", "3p", "4p", "5p",
    "gasolina", "alcool", "diesel", "gnv",
}

COMMERCIAL_KEYWORDS = {"furgao", "furgão", "ambulancia", "ambulância", "pick-up", "pickup", "cabine"}

PREMIUM_KEYWORDS = {"gti", "gts", "gtx", "turbo", "rally", "ss", "r-line", "rline", "cupê", "coupe",
                     "cabriolet", "conversivel", "conversível"}

TRIM_WORDS = {"gl", "gls", "glx", "ex", "exs", "lx", "lxs", "lt", "ltz", "ls", "se", "sx",
              "dx", "xs", "xr", "cl", "cx", "cd", "ghia", "life", "spirit", "maxx",
              "city", "trend", "comfort", "comfortline", "trendline", "highline",
              "attractive", "attractiv", "hatch", "sedan", "sed"}


def load_cache():
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_cache(cache):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


API_BLOCKED = False

def api_get(url, cache, retries=2, delay=2.0):
    global API_BLOCKED
    if url in cache:
        return cache[url]
    if API_BLOCKED:
        return None
    for attempt in range(retries):
        try:
            time.sleep(delay)
            resp = requests.get(url, timeout=15)
            if resp.status_code == 200:
                data = resp.json()
                if isinstance(data, dict) and "error" in data:
                    if "limite de taxa" in str(data.get("error", "")):
                        API_BLOCKED = True
                        print("    [API blocked by rate limit - using cache only]", flush=True)
                    return None
                cache[url] = data
                return data
            elif resp.status_code == 429:
                API_BLOCKED = True
                print("    [API rate limited - using cache only]", flush=True)
                return None
            else:
                return None
        except Exception:
            if attempt < retries - 1:
                time.sleep(3)
    return None


SHORT_CODES = {"C3", "C4", "C5", "C8", "X1", "X2", "X3", "X4", "X5", "X6", "S2", "S3", "S4", "S5", "A3", "A4", "A5", "A6", "Q3", "Q5", "Q7"}

def normalize_model(text):
    t = text.strip()
    t = re.sub(r'([A-Za-z])(\d)', r'\1 \2', t)
    t = re.sub(r'(\d)([A-Za-z])', r'\1 \2', t)
    t = re.sub(r'\s+', ' ', t).strip()
    for code in SHORT_CODES:
        spaced = f"{code[0]} {code[1:]}"
        t = re.sub(r'(?<![A-Za-z])' + re.escape(spaced) + r'(?![A-Za-z0-9])', code, t, flags=re.IGNORECASE)
    return t


def extract_cc(text):
    t = normalize_model(text).lower()
    matches = re.findall(r'\b(\d{2,4})\b', t)
    for m in matches:
        val = int(m)
        if 50 <= val <= 1000 and val not in range(1980, 2030):
            return val
    return None


def extract_displacement_liters(text):
    t = text.lower().replace(",", ".")
    m = re.search(r'(\d)\.(\d)\s*l?\b', t)
    if m:
        return float(f"{m.group(1)}.{m.group(2)}")
    m = re.search(r'\b(\d)(\d)l\b', t)
    if m:
        return float(f"{m.group(1)}.{m.group(2)}")
    m = re.search(r'\b(\d\.\d)\b(?!\s*v)', t)
    if m:
        return float(m.group(1))
    m = re.search(r'\b(1000|1500|1600|1800|2000)\b', t)
    if m:
        return int(m.group(1)) / 1000.0
    return None


def get_first_model_word(text):
    """Get the first significant word of the model (e.g., GOL, SCENIC, CELTA, FAZER, 206)."""
    norm = normalize_model(text).upper()
    words = norm.split()
    for w in words:
        if w.lower() not in NOISE_WORDS and len(w) >= 2:
            return w
    return words[0] if words else ""


def parse_vehicle(description, fab_mod):
    desc = description.strip()
    if desc.startswith("I/"):
        desc = desc[2:]
    elif desc.startswith("IMP/"):
        desc = desc[4:]

    if "/" in desc:
        brand = desc.split("/")[0].strip().upper()
        model = desc.split("/", 1)[1].strip()
    else:
        parts = desc.split()
        brand = parts[0].upper()
        model = " ".join(parts[1:]) if len(parts) > 1 else ""

    model_words = model.split()
    expanded = []
    for w in model_words:
        alias = MODEL_ALIASES.get(w.upper())
        expanded.append(alias if alias else w)
    model = " ".join(expanded)

    year = None
    if fab_mod:
        parts = str(fab_mod).strip().split("/")
        try:
            year = int(parts[-1].strip())
        except ValueError:
            pass

    return brand, model, year


def is_motorcycle(brand, model):
    brand_up = brand.upper()
    model_up = model.upper()
    if any(kw in model_up for kw in CAR_KEYWORDS):
        return False
    if brand_up in BRAND_MAP_MOTOS_API and brand_up not in BRAND_MAP_CARROS:
        return True
    if brand_up == "HONDA":
        return any(kw in model_up or kw in f" {model_up} " for kw in MOTO_KEYWORDS)
    return False


# ─── CSV-based matching for cars ─────────────────────────────────────────

def load_fipe_csv(filepath):
    """Load the FIPE CSV into a dictionary indexed by brand."""
    data = defaultdict(list)
    with open(filepath, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader:
            brand = row["Marca"]
            data[brand].append({
                "modelo": row["Modelo"],
                "ano": int(row["Ano"]) if row["Ano"].isdigit() else None,
                "valor": row["Valor"],
                "codigo_fipe": row["CodigoFipe"],
                "combustivel": row["Combustivel"],
            })
    return data


def get_trim_words(text):
    """Extract trim-level words from a model string."""
    words = set()
    cleaned = re.sub(r'[/\-\(\)]', ' ', normalize_model(text).lower())
    for w in cleaned.split():
        w = w.strip('.,;:')
        if w in TRIM_WORDS:
            words.add(w)
    return words


def score_model(model_clean, modelo, model_disp, input_trims, year, records):
    """Score a FIPE model candidate considering text similarity, trim match, year coverage, and commercial penalty."""
    fipe_lower = modelo.lower()
    fipe_lower_nopunct = re.sub(r'[/\-\(\)]', ' ', fipe_lower)

    # Reject commercial variants unless input explicitly mentions them
    if any(kw in fipe_lower_nopunct for kw in COMMERCIAL_KEYWORDS):
        if not any(kw in model_clean for kw in COMMERCIAL_KEYWORDS):
            return -1, None

    # Penalize premium/sporty variants unless input mentions them
    premium_penalty = 0
    for kw in PREMIUM_KEYWORDS:
        if kw in fipe_lower_nopunct.split() or kw in fipe_lower_nopunct:
            if kw not in model_clean:
                premium_penalty = -0.25
                break

    fipe_clean = re.sub(r'[^a-z0-9 ]', '', normalize_model(modelo).lower())
    text_score = SequenceMatcher(None, model_clean, fipe_clean).ratio() + premium_penalty

    # Trim bonus: reward matching trim words (GL, LX, Life, Spirit, etc.)
    fipe_trims = get_trim_words(modelo)
    trim_overlap = len(input_trims & fipe_trims)
    trim_penalty = len(input_trims - fipe_trims) * 0.05
    trim_bonus = trim_overlap * 0.15 - trim_penalty

    # Displacement bonus
    disp_bonus = 0
    if model_disp:
        fipe_disp = extract_displacement_liters(modelo)
        if fipe_disp:
            if abs(fipe_disp - model_disp) < 0.05:
                disp_bonus = 0.15
            elif abs(fipe_disp - model_disp) < 0.3:
                disp_bonus = 0.05
            else:
                disp_bonus = -0.10

    # Year coverage: prefer models whose year range covers the target year
    year_bonus = 0
    if year:
        model_records = [r for r in records if r["modelo"] == modelo]
        years = [r["ano"] for r in model_records if r["ano"]]
        if years:
            if year in years:
                year_bonus = 0.20
            else:
                min_diff = min(abs(y - year) for y in years)
                if min_diff <= 2:
                    year_bonus = 0.10
                elif min_diff <= 5:
                    year_bonus = 0.0
                else:
                    year_bonus = -0.15

    total = text_score + trim_bonus + disp_bonus + year_bonus

    # Find best year record for this model
    best_rec = None
    if year:
        model_records = [r for r in records if r["modelo"] == modelo]
        for r in model_records:
            if r["ano"] == year:
                best_rec = r
                break
        if not best_rec:
            closest_diff = 999
            for r in model_records:
                if r["ano"] and abs(r["ano"] - year) < closest_diff:
                    closest_diff = abs(r["ano"] - year)
                    best_rec = r
    if not best_rec:
        model_records = [r for r in records if r["modelo"] == modelo]
        if model_records:
            best_rec = model_records[0]

    return total, best_rec


def match_car_csv(brand_sheet, model_sheet, year, csv_data):
    """Match a car against the CSV dataset."""

    csv_brand = BRAND_MAP_CARROS.get(brand_sheet.upper())
    if not csv_brand:
        return None

    records = csv_data.get(csv_brand, [])
    if not records:
        for csv_b, recs in csv_data.items():
            if brand_sheet.upper() in csv_b.upper() or csv_b.upper() in brand_sheet.upper():
                records = recs
                break
    if not records:
        return None

    model_norm = normalize_model(model_sheet)
    model_first_word = get_first_model_word(model_sheet)
    model_disp = extract_displacement_liters(model_sheet)
    model_clean = re.sub(r'[^a-z0-9 ]', '', model_norm.lower())
    input_trims = get_trim_words(model_sheet)

    unique_models = {}
    for r in records:
        if r["modelo"] not in unique_models:
            unique_models[r["modelo"]] = r

    # Phase 1: Filter by core model name
    candidates = []
    for modelo in unique_models:
        fipe_first_word = get_first_model_word(modelo)
        if model_first_word == fipe_first_word:
            candidates.append(modelo)
        elif model_first_word in normalize_model(modelo).upper().split():
            candidates.append(modelo)

    if not candidates:
        for modelo in unique_models:
            if model_first_word in normalize_model(modelo).upper():
                candidates.append(modelo)

    if not candidates:
        return None

    # Phase 2: Score all candidates with combined metric
    best_total = -999
    best_modelo = None
    best_rec = None
    for modelo in candidates:
        total, rec = score_model(model_clean, modelo, model_disp, input_trims, year, records)
        if total > best_total and rec:
            best_total = total
            best_modelo = modelo
            best_rec = rec

    if not best_rec:
        return None

    return {
        "valor": best_rec["valor"],
        "modelo_fipe": best_modelo,
        "ano": best_rec["ano"],
        "score": best_total,
        "codigo_fipe": best_rec["codigo_fipe"],
    }


# ─── API-based matching for motos ────────────────────────────────────────

def match_moto_api(brand_sheet, model_sheet, year, cache):
    """Match a motorcycle against the FIPE API (cached)."""
    brand_code = BRAND_MAP_MOTOS_API.get(brand_sheet.upper())
    if not brand_code:
        return None

    # Get model list
    models_url = f"{API_BASE}/motos/marcas/{brand_code}/modelos"
    models_data = api_get(models_url, cache, delay=3.0)
    if not models_data or "modelos" not in models_data:
        return None

    models_list = models_data["modelos"]
    model_norm = normalize_model(model_sheet)
    model_first = get_first_model_word(model_sheet)
    model_cc = extract_cc(model_sheet)
    model_clean = re.sub(r'[^a-z0-9 ]', '', model_norm.lower())

    # Phase 1: Core name filter
    candidates = []
    for item in models_list:
        fipe_first = get_first_model_word(item["nome"])
        fipe_norm = normalize_model(item["nome"]).upper()
        if model_first == fipe_first:
            candidates.append(item)
        elif model_first in fipe_norm.split():
            candidates.append(item)

    if not candidates:
        for item in models_list:
            if model_first in normalize_model(item["nome"]).upper():
                candidates.append(item)

    if not candidates:
        return None

    # Phase 2: CC filter
    if model_cc:
        cc_exact = [it for it in candidates if extract_cc(it["nome"]) and abs(extract_cc(it["nome"]) - model_cc) <= 5]
        cc_close = [it for it in candidates if extract_cc(it["nome"]) and abs(extract_cc(it["nome"]) - model_cc) <= 25]
        if cc_exact:
            candidates = cc_exact
        elif cc_close:
            candidates = cc_close

    # Phase 3: Text match
    best_score = 0
    best_item = None
    for item in candidates:
        fipe_clean = re.sub(r'[^a-z0-9 ]', '', normalize_model(item["nome"]).lower())
        score = SequenceMatcher(None, model_clean, fipe_clean).ratio()
        if score > best_score:
            best_score = score
            best_item = item

    if not best_item:
        return None

    # Get year and price
    anos_url = f"{API_BASE}/motos/marcas/{brand_code}/modelos/{best_item['codigo']}/anos"
    anos_data = api_get(anos_url, cache, delay=1.5)
    if not anos_data:
        return None

    year_code = None
    if year:
        for ano in anos_data:
            try:
                if int(ano["nome"].split()[0]) == year:
                    year_code = ano["codigo"]
                    break
            except ValueError:
                continue
        if not year_code:
            closest_diff = 999
            closest = None
            for ano in anos_data:
                try:
                    diff = abs(int(ano["nome"].split()[0]) - year)
                    if diff < closest_diff:
                        closest_diff = diff
                        closest = ano
                except ValueError:
                    continue
            if closest and closest_diff <= 2:
                year_code = closest["codigo"]

    if not year_code and anos_data:
        year_code = anos_data[0]["codigo"]

    if not year_code:
        return None

    price_url = f"{API_BASE}/motos/marcas/{brand_code}/modelos/{best_item['codigo']}/anos/{year_code}"
    price_data = api_get(price_url, cache, delay=1.5)
    if price_data and "Valor" in price_data:
        return {
            "valor": price_data["Valor"],
            "modelo_fipe": best_item["nome"],
            "ano": price_data.get("AnoModelo"),
            "score": best_score,
            "codigo_fipe": price_data.get("CodigoFipe", ""),
        }
    return None


def main():
    if len(sys.argv) < 2:
        print("Usage: python fipe_enricher.py <input_file.xlsx>")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = f"{os.path.splitext(input_file)[0]}_com_fipe.xlsx"

    print(f"Reading {input_file}...", flush=True)
    wb = openpyxl.load_workbook(input_file)
    ws = wb.active
    total_rows = ws.max_row - 1

    # Load data sources
    print(f"Loading FIPE CSV ({CSV_FILE})...", flush=True)
    csv_data = load_fipe_csv(CSV_FILE)
    print(f"  {sum(len(v) for v in csv_data.values())} car records from {len(csv_data)} brands", flush=True)

    cache = load_cache()
    print(f"  {len(cache)} cached API entries\n", flush=True)

    found = 0
    not_found = 0
    errors = []

    for row in range(2, ws.max_row + 1):
        desc = ws.cell(row, 2).value
        fab_mod = ws.cell(row, 4).value
        lote = ws.cell(row, 1).value

        if not desc:
            continue

        brand, model, year = parse_vehicle(desc, fab_mod)
        idx = row - 1
        progress = f"[{idx}/{total_rows}]"

        moto = is_motorcycle(brand, model)

        try:
            if moto:
                result = match_moto_api(brand, model, year, cache)
                if not result:
                    result = match_car_csv(brand, model, year, csv_data)
                    if result:
                        moto = False
            else:
                result = match_car_csv(brand, model, year, csv_data)
                if not result and brand.upper() in BRAND_MAP_MOTOS_API:
                    result = match_moto_api(brand, model, year, cache)
                    if result:
                        moto = True

            if result:
                ws.cell(row, 6).value = result["valor"]
                found += 1
                vtype = "MOTO" if moto else "CAR"
                print(f"{progress} Lote {lote}: {brand}/{model} {year} -> {result['valor']} ({vtype}: {result['modelo_fipe']}, score: {result['score']:.2f})", flush=True)
            else:
                ws.cell(row, 6).value = "NAO ENCONTRADO"
                not_found += 1
                print(f"{progress} Lote {lote}: {desc} -> NOT FOUND", flush=True)
                errors.append(f"Lote {lote}: {desc} ({year})")

        except Exception as e:
            ws.cell(row, 6).value = "ERRO"
            not_found += 1
            print(f"{progress} Lote {lote}: ERROR - {e}", flush=True)
            errors.append(f"Lote {lote}: {desc} - {e}")

        if idx % 50 == 0:
            save_cache(cache)

    save_cache(cache)
    wb.save(output_file)

    print(f"\n{'='*50}")
    print(f"RESULTS:")
    print(f"  Total vehicles: {total_rows}")
    print(f"  Found: {found}")
    print(f"  Not found: {not_found}")
    print(f"  Success rate: {found/max(total_rows,1)*100:.1f}%")
    print(f"{'='*50}")

    if errors:
        print(f"\nNot found ({len(errors)}):")
        for e in errors:
            print(f"  - {e}")

    print(f"\nSaved to: {output_file}", flush=True)


if __name__ == "__main__":
    main()
